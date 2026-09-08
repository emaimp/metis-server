from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader

from app.core.settings import MAX_CHARS
from app.repositories.chats import get_attachment_db

MAX_PDF_PAGES = 30


def _truncate(content: str) -> str:
    if len(content) > MAX_CHARS:
        content = content[:MAX_CHARS] + (
            f"\n\n[... content truncated, only the first {MAX_CHARS} characters are shown ...]"
        )
    return content


def make_attachment_tools(chat_id: str) -> dict:
    """Build request-scoped read tools that load files from this chat's attachments in the database.

    Each tool takes the `attachment_id` given to the model in the hint and reads the raw
    bytes from the database (no temp files involved).
    """

    def _load(attachment_id: str) -> tuple[str, bytes] | str:
        attachment = get_attachment_db(chat_id, attachment_id)
        if attachment is None:
            return f"Error: The attachment '{attachment_id}' does not exist."
        return attachment["filename"], attachment["data"]

    def read_txt(attachment_id: str) -> str:
        """Reads a text document (.txt) stored as a chat attachment.

        Args:
            attachment_id: The id of the attached text document.
        """
        loaded = _load(attachment_id)
        if isinstance(loaded, str):
            return loaded
        _filename, data = loaded
        try:
            content = data.decode("utf-8", errors="replace")
        except Exception as e:
            return f"Error reading the document: {e}"
        return _truncate(content)

    def read_pdf(attachment_id: str) -> str:
        """Reads a PDF document stored as a chat attachment.

        Args:
            attachment_id: The id of the attached PDF document.
        """
        loaded = _load(attachment_id)
        if isinstance(loaded, str):
            return loaded
        _filename, data = loaded
        try:
            reader = PdfReader(io.BytesIO(data))
            total_pages = len(reader.pages)
            parts = []
            for page in reader.pages[:MAX_PDF_PAGES]:
                text = page.extract_text() or ""
                parts.append(text.strip())
            content = "\n\n".join(p for p in parts if p)
            if total_pages > MAX_PDF_PAGES:
                content += (
                    f"\n\n[... document is longer, only the first {MAX_PDF_PAGES} "
                    f"pages of {total_pages} were read ...]"
                )
            if not content.strip():
                return (
                    "Error: Could not extract text from the PDF. "
                    "It may be a scanned document without a text layer."
                )
        except Exception as e:
            return f"Error reading the PDF: {e}"
        return _truncate(content)

    def read_docx(attachment_id: str) -> str:
        """Reads a Word document (.docx) stored as a chat attachment.

        Args:
            attachment_id: The id of the attached Word document.
        """
        loaded = _load(attachment_id)
        if isinstance(loaded, str):
            return loaded
        _filename, data = loaded
        try:
            doc = Document(io.BytesIO(data))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    parts.append(" | ".join(cells))
            content = "\n".join(parts)
            if not content.strip():
                return (
                    "Error: Could not extract text from the Word document. "
                    "It may be empty or contain only images."
                )
        except Exception as e:
            return f"Error reading the Word document: {e}"
        return _truncate(content)

    return {"read_txt": read_txt, "read_pdf": read_pdf, "read_docx": read_docx}