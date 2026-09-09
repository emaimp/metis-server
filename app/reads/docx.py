import io

from docx import Document

from app.core.settings import MAX_CHARS


def extract_docx_from_bytes(data: bytes) -> str:
    """Extract text from a .docx document stored as bytes (no disk I/O), including paragraphs and tables."""
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        return f"Error reading the Word document: {str(e)}"

    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(' | '.join(cells))
    content = '\n'.join(parts)

    if not content.strip():
        return (
            "Error: Could not extract text from the Word document. "
            "It may be empty or contain only images."
        )

    if len(content) > MAX_CHARS:
        content = (
            content[:MAX_CHARS]
            + f"\n\n[... content truncated, only the first "
              f"{MAX_CHARS} characters are shown ...]"
        )

    return content

