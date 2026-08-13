import os

from docx import Document

from app.core.settings import MAX_CHARS


def extract_docx_text(file_path: str) -> str:
    """
    Extracts text from a Word document (.docx), including paragraphs and tables.

    Args:
        file_path: The path to the .docx file.
    """
    doc = Document(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(' | '.join(cells))
    return '\n'.join(parts)


# Tool definition for .docx documents
def read_docx(file_path: str) -> str:
    """
    Reads and extracts the content of a Word document (.docx) for general analysis.

    Args:
        file_path: The path to the Word document to read.
    """
    if not os.path.exists(file_path):
        return f"Error: The file '{file_path}' does not exist."

    try:
        content = extract_docx_text(file_path)
    except Exception as e:
        return f"Error reading the Word document: {str(e)}"

    if not content.strip():
        return (
            f"Error: Could not extract text from the Word document "
            f"'{file_path}'. It may be empty or contain only images."
        )

    if len(content) > MAX_CHARS:
        content = (
            content[:MAX_CHARS]
            + f"\n\n[... content truncated, only the first "
              f"{MAX_CHARS} characters are shown ...]"
        )

    return content
