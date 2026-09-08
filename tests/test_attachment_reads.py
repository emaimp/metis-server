import io
import pathlib
import sys

ROOT = str(pathlib.Path(__file__).resolve().parents[1])
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import pytest

from app.core.database import init_db
from app.repositories import chats as repo
from app.reads.attachment_reads import make_attachment_tools


@pytest.fixture()
def chat_id(tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> str:
    monkeypatch.setattr('app.core.database.DB_PATH', tmp_path / 'test_chats.db')
    init_db()
    return repo.create_chat_db('test')['id']


def test_read_txt_from_database(chat_id):
    att = repo.create_attachment_db(chat_id, 'nota.txt', 'text/plain', b'primera linea\nsegunda linea')
    tools = make_attachment_tools(chat_id)
    assert tools['read_txt'](att['id']) == 'primera linea\nsegunda linea'


def test_read_txt_truncates_to_max_chars(chat_id, monkeypatch):
    monkeypatch.setattr('app.reads.attachment_reads.MAX_CHARS', 10)
    att = repo.create_attachment_db(chat_id, 'nota.txt', 'text/plain', b'x' * 50)
    tools = make_attachment_tools(chat_id)
    result = tools['read_txt'](att['id'])
    assert result.startswith('x' * 10)
    assert 'content truncated' in result


def test_read_txt_missing_attachment_returns_error(chat_id):
    tools = make_attachment_tools(chat_id)
    assert tools['read_txt']('no-existe').startswith('Error')


def test_read_docx_from_database(chat_id):
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph('Hola mundo docx')
    doc.save(buf)

    att = repo.create_attachment_db(
        chat_id,
        'doc.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        buf.getvalue(),
    )
    tools = make_attachment_tools(chat_id)
    assert tools['read_docx'](att['id']) == 'Hola mundo docx'


def test_read_docx_empty_returns_error(chat_id):
    from docx import Document

    buf = io.BytesIO()
    Document().save(buf)
    att = repo.create_attachment_db(
        chat_id,
        'vacio.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        buf.getvalue(),
    )
    tools = make_attachment_tools(chat_id)
    assert tools['read_docx'](att['id']).startswith('Error')


def _minimal_pdf(text: str) -> bytes:
    stream = f'BT /F1 24 Tf 72 720 Td ({text}) Tj ET'.encode('latin-1')
    out = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        None,
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n'.encode()
        if i == 4:
            out += b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream\nendobj\n'
        else:
            out += obj + b'\nendobj\n'
    xref_pos = len(out)
    out += b'xref\n0 6\n0000000000 65535 f \n'
    for off in offsets[1:]:
        out += f'{off:010d} 00000 n \n'.encode()
    out += b'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n' + str(xref_pos).encode() + b'\n%%EOF\n'
    return bytes(out)


def test_read_pdf_from_database(chat_id):
    att = repo.create_attachment_db(chat_id, 'doc.pdf', 'application/pdf', _minimal_pdf('Hello DB'))
    tools = make_attachment_tools(chat_id)
    result = tools['read_pdf'](att['id'])
    assert 'Hello DB' in result